from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _style_header(worksheet):
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    worksheet.freeze_panes = "A2"

def create_workbook():
    wb = Workbook()

    # Main Report Sheet
    ws = wb.active
    ws.title = "Scam Reports"

    ws.append([
        "Subreddit",
        "Date of Post",
        "Title",
        "Summary",
        "Phone Numbers",
        "UPI IDs",
        "Account Numbers",
        "Emails",
        "Scam Type",
        "AI Confidence",
        "Reddit URL"
    ])
    _style_header(ws)

    # Summary Sheet
    summary_ws = wb.create_sheet("Summary")

    summary_ws.append([
        "Subreddit",
        "RSS Fetched",
        "Failed Keyword",
        "Passed Keyword",
        "No Contact Info",
        "Sent To AI",
        "AI Rejected",
        "Final Cases"
    ])
    _style_header(summary_ws)

    return wb, ws, summary_ws
  
def write_post(ws, record):
    ws.append([
        record.subreddit,
        record.post_date,
        record.title,
        record.summary[:500],
        ", ".join(record.phones),
        ", ".join(record.upi_ids),
        ", ".join(record.account_numbers),
        ", ".join(record.emails),
        record.scam_type,
        record.confidence,
        record.link
    ])   

def write_summary(
    summary_ws,
    subreddit,
    total_fetched,
    passed_keyword,
    sent_to_ai,
    passed_ai,
    final_with_contact_info,
):
    summary_ws.append([
        subreddit,
        total_fetched,
        total_fetched - passed_keyword,
        passed_keyword,
        passed_keyword - sent_to_ai,
        sent_to_ai,
        sent_to_ai - passed_ai,
        final_with_contact_info
    ])

def save_workbook(wb, filename):
    wb.save(filename)
